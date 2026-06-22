from __future__ import annotations

import os
from io import BytesIO
from threading import Lock
from typing import Any

from langchain_chroma import Chroma
from langchain_community.document_loaders.parsers.pdf import PyPDFium2Parser
from langchain_core.documents import Document
from langchain_core.documents.base import Blob
from langchain_core.prompts import ChatPromptTemplate
from langchain_groq import ChatGroq
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import CharacterTextSplitter
from docx import Document as DocxDocument

# ---------------------------------------------------------------------------
# Globals
# ---------------------------------------------------------------------------

_lock = Lock()

_vector_store: Chroma | None = None

_api_key = os.getenv("GROQ_API_KEY")
if not _api_key:
    raise RuntimeError("GROQ_API_KEY is not set")

_llm = ChatGroq(model="llama-3.1-8b-instant", groq_api_key=_api_key, temperature=0.2)

# ---------------------------------------------------------------------------
# Vector store (lazy-init to avoid uvicorn reload issues)
# ---------------------------------------------------------------------------


def _get_vector_store() -> Chroma:
    global _vector_store
    if _vector_store is None:
        _vector_store = Chroma(
            collection_name="documents",
            embedding_function=HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2"),
            persist_directory="./chroma_db",
            collection_metadata={"hnsw:space": "cosine"},
        )
    return _vector_store


# ---------------------------------------------------------------------------
# Document loading
# ---------------------------------------------------------------------------

_text_splitter = CharacterTextSplitter(separator="", chunk_size=500, chunk_overlap=0)


def _load_pdf(data: bytes, filename: str) -> list[Document]:
    blob = Blob.from_data(data, path=filename)
    return list(PyPDFium2Parser().lazy_parse(blob))


def _load_docx(data: bytes) -> list[Document]:
    doc = DocxDocument(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs).strip()
    return [Document(page_content=text)] if text else []


def _load_documents(filename: str, data: bytes) -> list[Document]:
    """Parse raw bytes into langchain Documents based on file extension."""
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return _load_pdf(data, filename)
    if ext == "docx":
        return _load_docx(data)
    raise ValueError("Only PDF and DOCX supported")


# ---------------------------------------------------------------------------
# Ingestion: load → chunk → store
# ---------------------------------------------------------------------------

def ingest(filename: str, data: bytes) -> int:
    """Ingest a file: extract text, chunk it, store in vector DB. Returns chunk count."""
    docs = _load_documents(filename, data)
    if not docs:
        return 0

    chunks = _text_splitter.split_documents(docs)
    doc_type = filename.rsplit(".", 1)[-1].lower()

    documents: list[Document] = []
    ids: list[str] = []

    for i, chunk in enumerate(chunks):
        text = (chunk.page_content or "").strip()
        if not text:
            continue
        meta = {"source": filename, "type": doc_type, "chunkindex": i}
        if doc_type == "pdf":
            meta["page"] = chunk.metadata.get("page")
        documents.append(Document(page_content=text, metadata=meta))
        ids.append(f"{filename}-{i}")

    if not documents:
        return 0

    with _lock:
        _get_vector_store().add_documents(documents, ids=ids)

    return len(documents)


# ---------------------------------------------------------------------------
# Search
# ---------------------------------------------------------------------------

def search(query: str, topk: int = 3, threshold: float = 0.4) -> list[dict[str, Any]]:
    """Similarity search. Returns list of {text, score, source, type, chunkindex, page}."""
    if not query or not query.strip():
        return []

    with _lock:
        results = _get_vector_store().similarity_search_with_score(query, k=topk)

    return [
        {
            "text": doc.page_content,
            "score": round(1 - float(distance), 4) if distance is not None else None,
            "source": (doc.metadata or {}).get("source"),
            "type": (doc.metadata or {}).get("type"),
            "chunkindex": (doc.metadata or {}).get("chunkindex"),
            "page": (doc.metadata or {}).get("page"),
        }
        for doc, distance in results
        if (1 - float(distance)) >= threshold
    ]


# ---------------------------------------------------------------------------
# RAG
# ---------------------------------------------------------------------------

RAG_PROMPT = ChatPromptTemplate.from_messages([
    ("system", "Answer based on the context provided. If not found, say you don't know."),
    ("human", "Context:\n{context}\n\nQuestion: {question}"),
])


def ask(question: str, threshold: float = 0.4) -> dict:
    """Simple RAG: retrieve relevant chunks, then ask the LLM."""
    if not question or not question.strip():
        return {"question": question, "answer": "", "sources": []}

    sources = search(question, topk=3, threshold=threshold)
    context = "\n\n".join(s["text"] for s in sources)

    chain = RAG_PROMPT | _llm
    answer = chain.invoke({"context": context, "question": question}).content

    return {
        "question": question,
        "answer": answer or "No answer generated.",
        "sources": sources,
    }
