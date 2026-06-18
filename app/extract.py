"""Document text extraction for PDF and DOCX.

Clean, focused functions. No LangChain.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import BinaryIO

import pypdfium2 as pdfium
from docx import Document


def extract_pdf_text(file: BinaryIO | BytesIO | bytes | Path | str) -> str:
    """Extract plain text from a PDF file.

    Accepts file-like object, bytes, Path, or str path.
    Returns concatenated text from all pages.
    """
    if isinstance(file, (str, Path)):
        pdf = pdfium.PdfDocument(str(file))
    elif isinstance(file, (bytes, bytearray)):
        pdf = pdfium.PdfDocument(BytesIO(file))
    else:
        # file-like
        file.seek(0)
        data = file.read()
        pdf = pdfium.PdfDocument(BytesIO(data))

    texts: list[str] = []
    for page in pdf:
        textpage = page.get_textpage()
        texts.append(textpage.get_text_bounded())

    return "\n".join(texts).strip()


def extract_docx_text(file: BinaryIO | BytesIO | bytes | Path | str) -> str:
    """Extract plain text from a DOCX file.

    Accepts file-like object, bytes, Path, or str path.
    Returns all paragraph text joined by newlines.
    """
    if isinstance(file, (str, Path)):
        doc = Document(str(file))
    elif isinstance(file, (bytes, bytearray)):
        doc = Document(BytesIO(file))
    else:
        file.seek(0)
        data = file.read()
        doc = Document(BytesIO(data))

    return "\n".join(p.text for p in doc.paragraphs).strip()


def extract_text(filename: str, file: BinaryIO | BytesIO | bytes) -> str:
    """Dispatch to the correct extractor based on filename extension.

    Raises ValueError for unsupported types.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_pdf_text(file)
    if lower.endswith(".docx"):
        return extract_docx_text(file)
    raise ValueError(f"Unsupported file type: {filename}. Only .pdf and .docx allowed.")
